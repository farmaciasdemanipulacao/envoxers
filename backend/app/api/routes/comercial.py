from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import and_, delete, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_envoxer, get_current_gestor_ou_admin
from app.core.config import settings
from app.db.session import get_db
from app.models.comercial import (
    ComercialActivity,
    ComercialAIGeneration,
    ComercialAudit,
    ComercialAuditSource,
    ComercialCadence,
    ComercialCadenceStep,
    ComercialInsight,
    ComercialIntegration,
    ComercialLead,
    ComercialLeadCadence,
    ComercialLeadChannel,
    ComercialLeadContact,
    ComercialLeadFact,
    ComercialLeadTag,
    ComercialMessage,
    ComercialOpportunity,
    ComercialPipelineStage,
    ComercialSuppression,
    ComercialTag,
    ComercialTask,
)
from app.models.envoxer import Envoxer
from app.models.servico import Servico
from app.services.comercial_ai import PROMPT_VERSION, gerar_json

router = APIRouter(prefix="/comercial", tags=["comercial"])

ORIGENS = ["prospeccao_ativa", "google", "instagram", "indicacao", "evento", "lista", "importacao", "outro"]
CANAIS = ["instagram", "facebook", "tiktok", "site", "google_maps", "tripadvisor", "ifood", "rappi", "99food", "linktree", "cardapio_digital"]
CANAIS_MENSAGEM = ["instagram", "whatsapp", "email"]
TIPOS_MENSAGEM = ["abordagem", "resposta", "follow_up", "objecao", "conversa", "convite", "reuniao"]
OBJECOES = ["ja_tem_agencia", "nao_tem_interesse", "sem_orcamento", "manda_apresentacao", "falar_com_outra_pessoa", "esta_satisfeito", "agora_nao", "quanto_custa", "quem_e_voce", "outro"]
CATEGORIAS_OURO = ["google", "instagram", "cardapio", "delivery", "reputacao", "posicionamento", "conversao", "reserva", "whatsapp", "seo", "site", "prova_social", "recorrencia", "ticket_medio", "experiencia", "outro"]
TIPOS_FATO = ["fato", "hipotese", "preferencia", "objecao", "pessoa", "fio_comercial"]
STATUS_FINAIS = {"ganho", "perdido", "sem_resposta", "nao_prospectar"}
STATUS_REUNIAO_OU_DEPOIS = {"reuniao_agendada", "proposta", "negociacao", "ganho"}
STATUS_ABORDADOS = {"mensagem_1_enviada", "respondeu", "em_conversa", "follow_up_1", "follow_up_2", "follow_up_3", "oportunidade", "reuniao_agendada", "proposta", "negociacao", "ganho", "perdido", "sem_resposta"}
STATUS_RESPONDERAM = {"respondeu", "em_conversa", "oportunidade", "reuniao_agendada", "proposta", "negociacao", "ganho", "perdido"}


def _value(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def d(obj):
    return {c.name: _value(getattr(obj, c.name)) for c in obj.__table__.columns}


def _now():
    return datetime.now(timezone.utc)


def _norm(v: Optional[str]) -> str:
    return re.sub(r"\s+", " ", (v or "").strip().lower())


def _phone(v: Optional[str]) -> str:
    return re.sub(r"\D", "", v or "")


def _parse_date(v):
    if not v:
        return None
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_datetime(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    s = str(v).strip()
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


async def lead_or_404(db: AsyncSession, lead_id: int) -> ComercialLead:
    lead = (await db.execute(
        select(ComercialLead).where(ComercialLead.id == lead_id, ComercialLead.deleted_at.is_(None))
    )).scalar_one_or_none()
    if lead is None:
        raise HTTPException(404, "Lead não encontrado")
    return lead


async def activity(db: AsyncSession, lead_id: int, user_id: Optional[int], tipo: str, descricao: str, dados=None):
    db.add(ComercialActivity(
        lead_id=lead_id,
        usuario_envoxer_id=user_id,
        tipo=tipo,
        descricao=descricao,
        dados=dados,
    ))


async def _responsaveis(db: AsyncSession):
    rows = (await db.execute(
        select(Envoxer).where(Envoxer.ativo.is_(True), Envoxer.deleted_at.is_(None)).order_by(Envoxer.nome)
    )).scalars().all()
    return {x.id: x for x in rows}


async def _tags_por_lead(db: AsyncSession, ids: list[int]):
    out = defaultdict(list)
    if not ids:
        return out
    rows = (await db.execute(
        select(ComercialLeadTag.lead_id, ComercialTag.id, ComercialTag.nome)
        .join(ComercialTag, ComercialTag.id == ComercialLeadTag.tag_id)
        .where(ComercialLeadTag.lead_id.in_(ids))
        .order_by(ComercialTag.nome)
    )).all()
    for lead_id, tag_id, nome in rows:
        out[lead_id].append({"id": tag_id, "nome": nome})
    return out


async def _lead_view(db: AsyncSession, lead: ComercialLead, responsaveis=None, tags=None):
    if responsaveis is None:
        responsaveis = await _responsaveis(db)
    if tags is None:
        tags = await _tags_por_lead(db, [lead.id])
    item = d(lead)
    resp = responsaveis.get(lead.responsavel_envoxer_id)
    item["responsavel_nome"] = resp.nome if resp else None
    item["responsavel_foto"] = resp.foto_url if resp else None
    item["tags"] = tags.get(lead.id, [])
    channels = (await db.execute(
        select(ComercialLeadChannel).where(
            ComercialLeadChannel.lead_id == lead.id,
            ComercialLeadChannel.ativo.is_(True),
        )
    )).scalars().all()
    for channel in channels:
        item[channel.tipo] = channel.valor
    return item


async def contexto(db: AsyncSession, lead: ComercialLead):
    channels = [d(x) for x in (await db.execute(
        select(ComercialLeadChannel).where(ComercialLeadChannel.lead_id == lead.id, ComercialLeadChannel.ativo.is_(True))
    )).scalars().all()]
    sources = [d(x) for x in (await db.execute(
        select(ComercialAuditSource).where(ComercialAuditSource.lead_id == lead.id).order_by(ComercialAuditSource.consultado_em.desc())
    )).scalars().all()]
    insights = [d(x) for x in (await db.execute(
        select(ComercialInsight).where(ComercialInsight.lead_id == lead.id, ComercialInsight.ativo.is_(True)).order_by(ComercialInsight.created_at.desc())
    )).scalars().all()]
    messages = [d(x) for x in (await db.execute(
        select(ComercialMessage).where(ComercialMessage.lead_id == lead.id).order_by(ComercialMessage.enviado_em)
    )).scalars().all()]
    contacts = [d(x) for x in (await db.execute(
        select(ComercialLeadContact).where(ComercialLeadContact.lead_id == lead.id).order_by(ComercialLeadContact.principal.desc(), ComercialLeadContact.created_at)
    )).scalars().all()]
    facts = [d(x) for x in (await db.execute(
        select(ComercialLeadFact).where(ComercialLeadFact.lead_id == lead.id, ComercialLeadFact.ativo.is_(True)).order_by(ComercialLeadFact.created_at)
    )).scalars().all()]
    return {
        "lead": d(lead),
        "canais": channels,
        "fontes": sources,
        "ouros": insights,
        "mensagens": messages,
        "contatos": contacts,
        "memoria": facts,
    }


async def _close_or_pause_cadence(db: AsyncSession, lead_id: int, action: str):
    now = _now()
    values = {"status": action}
    if action == "pausada":
        values["pausada_em"] = now
    else:
        values["encerrada_em"] = now
        values["proximo_step_em"] = None
    await db.execute(
        ComercialLeadCadence.__table__.update()
        .where(ComercialLeadCadence.lead_id == lead_id, ComercialLeadCadence.status == "ativa")
        .values(**values)
    )


async def _status_side_effects(db: AsyncSession, lead: ComercialLead, old: str, new: str):
    if new in STATUS_REUNIAO_OU_DEPOIS:
        await _close_or_pause_cadence(db, lead.id, "encerrada")
    elif new == "pausado":
        await _close_or_pause_cadence(db, lead.id, "pausada")
    elif new in {"perdido", "sem_resposta"}:
        await _close_or_pause_cadence(db, lead.id, "encerrada")
    if new == "nao_prospectar":
        lead.nao_prospectar = True
        lead.proxima_acao = None
        lead.proxima_acao_em = None
        await _close_or_pause_cadence(db, lead.id, "encerrada")


async def _advance_active_cadence(db: AsyncSession, lead: ComercialLead):
    lc = (await db.execute(
        select(ComercialLeadCadence)
        .where(ComercialLeadCadence.lead_id == lead.id, ComercialLeadCadence.status == "ativa")
        .order_by(ComercialLeadCadence.created_at.desc())
    )).scalars().first()
    if not lc:
        return None
    current_order = lc.proximo_step_ordem
    current = (await db.execute(
        select(ComercialCadenceStep).where(
            ComercialCadenceStep.cadence_id == lc.cadence_id,
            ComercialCadenceStep.ordem == current_order,
        )
    )).scalar_one_or_none()
    nxt = (await db.execute(
        select(ComercialCadenceStep)
        .where(ComercialCadenceStep.cadence_id == lc.cadence_id, ComercialCadenceStep.ordem > current_order)
        .order_by(ComercialCadenceStep.ordem)
    )).scalars().first()
    if current:
        if current.ordem == 1:
            lead.status_codigo = "mensagem_1_enviada"
        elif current.ordem == 2:
            lead.status_codigo = "follow_up_1"
        elif current.ordem == 3:
            lead.status_codigo = "follow_up_2"
        elif current.ordem >= 4:
            lead.status_codigo = "follow_up_3"
    if nxt:
        lc.proximo_step_ordem = nxt.ordem
        lc.proximo_step_em = lc.iniciada_em + timedelta(days=nxt.dias_apos_inicio)
        lead.proxima_acao = {
            "follow_up": "Enviar follow-up com nova pepita",
            "reativacao": "Reativar com novo contexto",
            "abordagem": "Enviar abordagem",
        }.get(nxt.tipo, "Executar próxima etapa da cadência")
        lead.proxima_acao_em = lc.proximo_step_em
    else:
        lc.status = "concluida"
        lc.encerrada_em = _now()
        lc.proximo_step_em = None
        lead.proxima_acao = "Revisar lead após fim da cadência"
        lead.proxima_acao_em = _now() + timedelta(days=7)
    return current


async def _suppression_add(db: AsyncSession, lead: ComercialLead, user_id: int, motivo: str):
    values = [("email", lead.email), ("whatsapp", lead.whatsapp), ("telefone", lead.telefone)]
    channels = (await db.execute(
        select(ComercialLeadChannel).where(ComercialLeadChannel.lead_id == lead.id, ComercialLeadChannel.ativo.is_(True))
    )).scalars().all()
    for ch in channels:
        if ch.tipo in ("instagram", "facebook", "tiktok"):
            values.append((ch.tipo, ch.valor))
    for tipo, valor in values:
        if not valor:
            continue
        exists = (await db.execute(
            select(ComercialSuppression).where(
                ComercialSuppression.tipo == tipo,
                func.lower(ComercialSuppression.valor) == valor.lower(),
            )
        )).scalar_one_or_none()
        if not exists:
            db.add(ComercialSuppression(
                lead_id=lead.id,
                tipo=tipo,
                valor=valor,
                motivo=motivo,
                created_by_envoxer_id=user_id,
            ))


@router.get("/meta")
async def meta(
    db: Annotated[AsyncSession, Depends(get_db)],
    user: Annotated[Envoxer, Depends(get_current_envoxer)],
):
    stages = [d(x) for x in (await db.execute(
        select(ComercialPipelineStage).where(ComercialPipelineStage.ativo.is_(True)).order_by(ComercialPipelineStage.ordem)
    )).scalars().all()]
    users = [
        {"id": x.id, "nome": x.nome, "cargo": x.cargo, "permissao": x.permissao, "foto_url": x.foto_url}
        for x in (await db.execute(
            select(Envoxer).where(Envoxer.ativo.is_(True), Envoxer.deleted_at.is_(None)).order_by(Envoxer.nome)
        )).scalars().all()
    ]
    services = [{"id": x.id, "nome": x.nome, "slug": x.slug} for x in (await db.execute(select(Servico).where(Servico.ativo.is_(True)).order_by(Servico.nome))).scalars().all()]
    tags = [d(x) for x in (await db.execute(select(ComercialTag).order_by(ComercialTag.nome))).scalars().all()]
    return {
        "stages": stages,
        "users": users,
        "services": services,
        "tags": tags,
        "origens": ORIGENS,
        "canais": CANAIS,
        "canais_mensagem": CANAIS_MENSAGEM,
        "tipos_mensagem": TIPOS_MENSAGEM,
        "objecoes": OBJECOES,
        "categorias_ouro": CATEGORIAS_OURO,
        "tipos_fato": TIPOS_FATO,
        "permission": user.permissao,
    }


@router.get("/stages")
async def stages(
    db: Annotated[AsyncSession, Depends(get_db)],
    _: Annotated[Envoxer, Depends(get_current_envoxer)],
):
    return [d(x) for x in (await db.execute(
        select(ComercialPipelineStage).where(ComercialPipelineStage.ativo.is_(True)).order_by(ComercialPipelineStage.ordem)
    )).scalars().all()]


@router.post("/stages", status_code=201)
async def criar_stage(
    payload: dict,
    db: Annotated[AsyncSession, Depends(get_db)],
    _: Annotated[Envoxer, Depends(get_current_gestor_ou_admin)],
):
    nome = (payload.get("nome") or "").strip()
    codigo = _norm(payload.get("codigo") or nome).replace(" ", "_")
    codigo = re.sub(r"[^a-z0-9_]+", "", codigo)
    if not nome or not codigo:
        raise HTTPException(400, "Nome do status é obrigatório")
    x = ComercialPipelineStage(codigo=codigo, nome=nome, ordem=int(payload.get("ordem", 999)), tipo=payload.get("tipo", "aberto"), ativo=True)
    db.add(x)
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(409, "Já existe um status com esse código")
    await db.refresh(x)
    return d(x)


@router.patch("/stages/{stage_id}")
async def editar_stage(
    stage_id: int,
    payload: dict,
    db: Annotated[AsyncSession, Depends(get_db)],
    _: Annotated[Envoxer, Depends(get_current_gestor_ou_admin)],
):
    x = (await db.execute(select(ComercialPipelineStage).where(ComercialPipelineStage.id == stage_id))).scalar_one_or_none()
    if not x:
        raise HTTPException(404, "Status não encontrado")
    for key in ("nome", "ordem", "tipo", "ativo"):
        if key in payload:
            setattr(x, key, payload[key])
    await db.commit()
    await db.refresh(x)
    return d(x)


@router.get("/dashboard")
async def dashboard(
    db: Annotated[AsyncSession, Depends(get_db)],
    user: Annotated[Envoxer, Depends(get_current_envoxer)],
):
    q = select(ComercialLead).where(ComercialLead.deleted_at.is_(None))
    if user.permissao == "envoxer":
        q = q.where(or_(ComercialLead.responsavel_envoxer_id == user.id, ComercialLead.responsavel_envoxer_id.is_(None)))
    leads = (await db.execute(q)).scalars().all()
    ids = [x.id for x in leads]
    msgs = (await db.execute(select(ComercialMessage).where(ComercialMessage.lead_id.in_(ids)))).scalars().all() if ids else []
    audits = (await db.execute(select(ComercialAudit.lead_id).where(ComercialAudit.lead_id.in_(ids)).distinct())).all() if ids else []
    sent_leads = {m.lead_id for m in msgs if m.direcao == "enviada"}
    recv_leads = {m.lead_id for m in msgs if m.direcao == "recebida"}
    today = date.today()
    month_start = today.replace(day=1)
    won = [x for x in leads if x.status_codigo == "ganho"]
    funnel = [
        {"key": "leads", "label": "Leads", "value": len(leads)},
        {"key": "abordados", "label": "Abordados", "value": len(sent_leads)},
        {"key": "responderam", "label": "Responderam", "value": len(recv_leads)},
        {"key": "conversa", "label": "Conversa", "value": sum(x.status_codigo in STATUS_RESPONDERAM for x in leads)},
        {"key": "reuniao", "label": "Reunião", "value": sum(x.status_codigo in STATUS_REUNIAO_OU_DEPOIS for x in leads)},
        {"key": "proposta", "label": "Proposta", "value": sum(x.status_codigo in {"proposta", "negociacao", "ganho"} for x in leads)},
        {"key": "cliente", "label": "Cliente", "value": len(won)},
    ]
    top = sorted(leads, key=lambda x: (x.score_total or 0, x.updated_at or _now()), reverse=True)[:8]
    resp = await _responsaveis(db)
    tags = await _tags_por_lead(db, [x.id for x in top])
    return {
        "leads": len(leads),
        "novos": sum(x.data_entrada and x.data_entrada >= month_start for x in leads),
        "analisados": len({r[0] for r in audits}),
        "mensagens_enviadas": sum(m.direcao == "enviada" for m in msgs),
        "respostas": sum(m.direcao == "recebida" for m in msgs),
        "taxa_resposta": round((len(recv_leads) / len(sent_leads) * 100), 1) if sent_leads else 0,
        "conversas": sum(x.status_codigo in {"respondeu", "em_conversa"} for x in leads),
        "reunioes": sum(x.status_codigo == "reuniao_agendada" for x in leads),
        "propostas": sum(x.status_codigo == "proposta" for x in leads),
        "ganhos": len(won),
        "perdidos": sum(x.status_codigo == "perdido" for x in leads),
        "funnel": funnel,
        "top_leads": [await _lead_view(db, x, resp, tags) for x in top],
    }


@router.get("/hoje")
async def hoje(
    db: Annotated[AsyncSession, Depends(get_db)],
    user: Annotated[Envoxer, Depends(get_current_envoxer)],
):
    now = _now()
    base = select(ComercialLead).where(ComercialLead.deleted_at.is_(None), ComercialLead.nao_prospectar.is_(False))
    if user.permissao == "envoxer":
        base = base.where(or_(ComercialLead.responsavel_envoxer_id == user.id, ComercialLead.responsavel_envoxer_id.is_(None)))
    leads = (await db.execute(base)).scalars().all()
    ids = [x.id for x in leads]
    sent_ids = set()
    recv_ids = set()
    audit_ids = set()
    if ids:
        rows = (await db.execute(select(ComercialMessage.lead_id, ComercialMessage.direcao).where(ComercialMessage.lead_id.in_(ids)))).all()
        sent_ids = {lid for lid, direcao in rows if direcao == "enviada"}
        recv_ids = {lid for lid, direcao in rows if direcao == "recebida"}
        audit_ids = {r[0] for r in (await db.execute(select(ComercialAudit.lead_id).where(ComercialAudit.lead_id.in_(ids)).distinct())).all()}
    resp = await _responsaveis(db)
    tags = await _tags_por_lead(db, ids)

    async def items(seq):
        return [await _lead_view(db, x, resp, tags) for x in sorted(seq, key=lambda z: (z.proxima_acao_em or datetime.min.replace(tzinfo=timezone.utc), -(z.score_total or 0)))]

    # Seções mutuamente exclusivas: o mesmo lead não deve aparecer cinco vezes
    # na rotina diária. Priorizamos resposta > reunião > analisado > sem primeira
    # mensagem > ação vencida genérica.
    waiting = [x for x in leads if x.id in recv_ids and x.status_codigo == "respondeu"]
    used_ids = {x.id for x in waiting}
    meetings = [x for x in leads if x.id not in used_ids and x.status_codigo == "reuniao_agendada" and (not x.proxima_acao_em or x.proxima_acao_em.date() <= now.date())]
    used_ids.update(x.id for x in meetings)
    analyzed = [x for x in leads if x.id not in used_ids and x.id in audit_ids and x.id not in sent_ids]
    used_ids.update(x.id for x in analyzed)
    no_first = [x for x in leads if x.id not in used_ids and x.id not in sent_ids and x.status_codigo in {"novo", "a_pesquisar", "analise_concluida", "pronto_abordagem"}]
    used_ids.update(x.id for x in no_first)
    overdue = [x for x in leads if x.id not in used_ids and x.proxima_acao_em and x.proxima_acao_em <= now and x.status_codigo not in STATUS_FINAIS]

    tq = select(ComercialTask, ComercialLead).join(ComercialLead, ComercialLead.id == ComercialTask.lead_id).where(
        ComercialTask.status == "pendente",
        ComercialTask.prazo.is_not(None),
        ComercialTask.prazo <= now,
        ComercialLead.deleted_at.is_(None),
    )
    if user.permissao == "envoxer":
        tq = tq.where(or_(ComercialTask.responsavel_envoxer_id == user.id, ComercialTask.responsavel_envoxer_id.is_(None)))
    task_rows = (await db.execute(tq.order_by(ComercialTask.prazo))).all()
    overdue_tasks = [{**d(t), "lead": await _lead_view(db, l, resp, tags)} for t, l in task_rows]
    sections = [
        {"key": "respostas", "label": "Respostas aguardando ação", "items": await items(waiting)},
        {"key": "vencidos", "label": "Próximas ações vencidas", "items": await items(overdue)},
        {"key": "sem_primeira", "label": "Leads sem primeira mensagem", "items": await items(no_first)},
        {"key": "analisados", "label": "Analisados sem abordagem", "items": await items(analyzed)},
        {"key": "reunioes", "label": "Reuniões / ações de reunião", "items": await items(meetings)},
    ]
    unique = {x.id for seq in (waiting, overdue, no_first, analyzed, meetings) for x in seq}
    return {"total": len(unique) + len(overdue_tasks), "sections": sections, "tarefas_vencidas": overdue_tasks}


@router.get("/leads")
async def listar_leads(
    q: Optional[str] = None,
    status: Optional[str] = None,
    segmento: Optional[str] = None,
    cidade: Optional[str] = None,
    responsavel: Optional[int] = None,
    origem: Optional[str] = None,
    min_score: Optional[int] = None,
    com_resposta: Optional[bool] = None,
    com_auditoria: Optional[bool] = None,
    tag: Optional[int] = None,
    data_de: Optional[date] = None,
    data_ate: Optional[date] = None,
    ultimo_contato_de: Optional[date] = None,
    ultimo_contato_ate: Optional[date] = None,
    proxima_acao_de: Optional[date] = None,
    proxima_acao_ate: Optional[date] = None,
    db: AsyncSession = Depends(get_db),
    user: Envoxer = Depends(get_current_envoxer),
):
    st = select(ComercialLead).where(ComercialLead.deleted_at.is_(None))
    if user.permissao == "envoxer":
        st = st.where(or_(ComercialLead.responsavel_envoxer_id == user.id, ComercialLead.responsavel_envoxer_id.is_(None)))
    if status:
        st = st.where(ComercialLead.status_codigo == status)
    if segmento:
        st = st.where(ComercialLead.segmento.ilike(f"%{segmento}%"))
    if cidade:
        st = st.where(ComercialLead.cidade.ilike(f"%{cidade}%"))
    if responsavel:
        st = st.where(ComercialLead.responsavel_envoxer_id == responsavel)
    if origem:
        st = st.where(ComercialLead.origem == origem)
    if min_score is not None:
        st = st.where(ComercialLead.score_total >= min_score)
    if data_de:
        st = st.where(ComercialLead.data_entrada >= data_de)
    if data_ate:
        st = st.where(ComercialLead.data_entrada <= data_ate)
    if ultimo_contato_de:
        st = st.where(func.date(ComercialLead.ultimo_contato_em) >= ultimo_contato_de)
    if ultimo_contato_ate:
        st = st.where(func.date(ComercialLead.ultimo_contato_em) <= ultimo_contato_ate)
    if proxima_acao_de:
        st = st.where(func.date(ComercialLead.proxima_acao_em) >= proxima_acao_de)
    if proxima_acao_ate:
        st = st.where(func.date(ComercialLead.proxima_acao_em) <= proxima_acao_ate)
    if q:
        like = f"%{q}%"
        st = st.where(or_(
            ComercialLead.nome_estabelecimento.ilike(like),
            ComercialLead.nome_fantasia.ilike(like),
            ComercialLead.segmento.ilike(like),
            ComercialLead.cidade.ilike(like),
            ComercialLead.email.ilike(like),
            ComercialLead.whatsapp.ilike(like),
        ))
    if com_resposta is not None:
        exists_response = select(ComercialMessage.id).where(
            ComercialMessage.lead_id == ComercialLead.id,
            ComercialMessage.direcao == "recebida",
        ).exists()
        st = st.where(exists_response if com_resposta else ~exists_response)
    if com_auditoria is not None:
        exists_audit = select(ComercialAudit.id).where(ComercialAudit.lead_id == ComercialLead.id).exists()
        st = st.where(exists_audit if com_auditoria else ~exists_audit)
    if tag:
        st = st.where(select(ComercialLeadTag.lead_id).where(
            ComercialLeadTag.lead_id == ComercialLead.id,
            ComercialLeadTag.tag_id == tag,
        ).exists())
    leads = (await db.execute(st.order_by(ComercialLead.updated_at.desc()))).scalars().all()
    resp = await _responsaveis(db)
    tags = await _tags_por_lead(db, [x.id for x in leads])
    return [await _lead_view(db, x, resp, tags) for x in leads]


@router.post("/leads", status_code=201)
async def criar_lead(
    payload: dict,
    db: AsyncSession = Depends(get_db),
    user: Envoxer = Depends(get_current_envoxer),
):
    channels = payload.pop("canais", {}) or {}
    contact = payload.pop("contato", None)
    tags = payload.pop("tags", []) or []
    allowed = {c.name for c in ComercialLead.__table__.columns} - {"id", "created_at", "updated_at", "deleted_at"}
    data = {k: v for k, v in payload.items() if k in allowed}
    if not (data.get("nome_estabelecimento") or "").strip():
        raise HTTPException(400, "Nome do estabelecimento é obrigatório")
    if data.get("data_entrada"):
        data["data_entrada"] = _parse_date(data["data_entrada"]) or date.today()
    if data.get("proxima_acao_em"):
        data["proxima_acao_em"] = _parse_datetime(data["proxima_acao_em"])
    if user.permissao == "envoxer" and not data.get("responsavel_envoxer_id"):
        data["responsavel_envoxer_id"] = user.id
    if not data.get("proxima_acao"):
        data["proxima_acao"] = "Pesquisar e analisar lead"
    if not data.get("proxima_acao_em"):
        data["proxima_acao_em"] = _now()
    lead = ComercialLead(**data)
    db.add(lead)
    await db.flush()
    for tipo, valor in channels.items():
        if tipo in CANAIS and str(valor or "").strip():
            db.add(ComercialLeadChannel(lead_id=lead.id, tipo=tipo, valor=str(valor).strip()))
    if contact and (contact.get("nome") or "").strip():
        db.add(ComercialLeadContact(
            lead_id=lead.id,
            principal=True,
            **{k: v for k, v in contact.items() if k in {"nome", "cargo", "telefone", "whatsapp", "email"}},
        ))
    for tag_name in tags:
        await _set_tag(db, lead.id, str(tag_name))
    await activity(db, lead.id, user.id, "lead_criado", "Lead criado")
    await db.commit()
    await db.refresh(lead)
    return await _lead_view(db, lead)


@router.get("/leads/{lead_id}")
async def detalhe_lead(
    lead_id: int,
    db: AsyncSession = Depends(get_db),
    _: Envoxer = Depends(get_current_envoxer),
):
    lead = await lead_or_404(db, lead_id)
    ctx = await contexto(db, lead)
    ctx["lead"] = await _lead_view(db, lead)
    ctx["auditorias"] = [d(x) for x in (await db.execute(
        select(ComercialAudit).where(ComercialAudit.lead_id == lead_id).order_by(ComercialAudit.created_at.desc())
    )).scalars().all()]
    ctx["historico"] = [d(x) for x in (await db.execute(
        select(ComercialActivity).where(ComercialActivity.lead_id == lead_id).order_by(ComercialActivity.created_at.desc()).limit(250)
    )).scalars().all()]
    ctx["tarefas"] = [d(x) for x in (await db.execute(
        select(ComercialTask).where(ComercialTask.lead_id == lead_id).order_by(ComercialTask.status, ComercialTask.prazo.asc().nullslast())
    )).scalars().all()]
    ctx["oportunidades"] = [d(x) for x in (await db.execute(
        select(ComercialOpportunity).where(ComercialOpportunity.lead_id == lead_id).order_by(ComercialOpportunity.created_at.desc())
    )).scalars().all()]
    cadences = (await db.execute(
        select(ComercialLeadCadence).where(ComercialLeadCadence.lead_id == lead_id).order_by(ComercialLeadCadence.created_at.desc())
    )).scalars().all()
    cadence_views = []
    for lc in cadences:
        c = (await db.execute(select(ComercialCadence).where(ComercialCadence.id == lc.cadence_id))).scalar_one_or_none()
        steps = [d(x) for x in (await db.execute(
            select(ComercialCadenceStep).where(ComercialCadenceStep.cadence_id == lc.cadence_id).order_by(ComercialCadenceStep.ordem)
        )).scalars().all()]
        cadence_views.append({**d(lc), "nome": c.nome if c else None, "steps": steps})
    ctx["cadencias"] = cadence_views
    ctx["ia"] = [d(x) for x in (await db.execute(
        select(ComercialAIGeneration).where(ComercialAIGeneration.lead_id == lead_id).order_by(ComercialAIGeneration.created_at.desc()).limit(20)
    )).scalars().all()]
    return ctx


@router.patch("/leads/{lead_id}")
async def editar_lead(
    lead_id: int,
    payload: dict,
    db: AsyncSession = Depends(get_db),
    user: Envoxer = Depends(get_current_envoxer),
):
    lead = await lead_or_404(db, lead_id)
    old = lead.status_codigo
    allowed = {c.name for c in ComercialLead.__table__.columns} - {"id", "created_at", "updated_at", "deleted_at", "tenant_id"}
    for k, v in payload.items():
        if k not in allowed:
            continue
        if k == "data_entrada":
            v = _parse_date(v)
        elif k in {"proxima_acao_em", "ultimo_contato_em"}:
            v = _parse_datetime(v)
        elif k in {"fit_score", "opportunity_score", "engagement_score", "score_total"} and v is not None:
            v = max(0, min(100, int(v)))
        setattr(lead, k, v)
    if any(k in payload for k in ("fit_score", "opportunity_score", "engagement_score")) and "score_total" not in payload:
        lead.score_total = round((lead.fit_score + lead.opportunity_score + lead.engagement_score) / 3)
    if old != lead.status_codigo:
        await _status_side_effects(db, lead, old, lead.status_codigo)
        await activity(db, lead.id, user.id, "status_alterado", f"Status: {old} → {lead.status_codigo}", {"antes": old, "depois": lead.status_codigo})
    else:
        await activity(db, lead.id, user.id, "lead_atualizado", "Dados do lead atualizados")
    await db.commit()
    await db.refresh(lead)
    return await _lead_view(db, lead)


@router.delete("/leads/{lead_id}", status_code=204)
async def arquivar_lead(
    lead_id: int,
    db: Annotated[AsyncSession, Depends(get_db)],
    user: Annotated[Envoxer, Depends(get_current_gestor_ou_admin)],
):
    lead = await lead_or_404(db, lead_id)
    lead.deleted_at = _now()
    await _close_or_pause_cadence(db, lead_id, "encerrada")
    await activity(db, lead_id, user.id, "lead_arquivado", "Lead arquivado")
    await db.commit()


@router.post("/leads/{lead_id}/contatos", status_code=201)
async def criar_contato(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    nome = (payload.get("nome") or "").strip()
    if not nome:
        raise HTTPException(400, "Nome do contato é obrigatório")
    if payload.get("principal"):
        await db.execute(ComercialLeadContact.__table__.update().where(ComercialLeadContact.lead_id == lead_id).values(principal=False))
    x = ComercialLeadContact(lead_id=lead_id, **{k: payload.get(k) for k in ("nome", "cargo", "telefone", "whatsapp", "email", "principal") if k in payload})
    db.add(x)
    await activity(db, lead_id, user.id, "contato_adicionado", f"Contato adicionado: {nome}")
    await db.commit(); await db.refresh(x)
    return d(x)


@router.patch("/leads/{lead_id}/contatos/{contato_id}")
async def editar_contato(lead_id: int, contato_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    x = (await db.execute(select(ComercialLeadContact).where(ComercialLeadContact.id == contato_id, ComercialLeadContact.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Contato não encontrado")
    if payload.get("principal"):
        await db.execute(ComercialLeadContact.__table__.update().where(ComercialLeadContact.lead_id == lead_id).values(principal=False))
    for k in ("nome", "cargo", "telefone", "whatsapp", "email", "principal"):
        if k in payload: setattr(x, k, payload[k])
    await activity(db, lead_id, user.id, "contato_atualizado", f"Contato atualizado: {x.nome}")
    await db.commit(); await db.refresh(x)
    return d(x)


@router.delete("/leads/{lead_id}/contatos/{contato_id}", status_code=204)
async def excluir_contato(lead_id: int, contato_id: int, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    x = (await db.execute(select(ComercialLeadContact).where(ComercialLeadContact.id == contato_id, ComercialLeadContact.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Contato não encontrado")
    await db.delete(x); await activity(db, lead_id, user.id, "contato_removido", "Contato removido"); await db.commit()


@router.post("/leads/{lead_id}/canais", status_code=201)
async def upsert_canal(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    tipo = payload.get("tipo")
    valor = (payload.get("valor") or "").strip()
    if tipo not in CANAIS or not valor:
        raise HTTPException(400, "Canal e valor são obrigatórios")
    x = (await db.execute(select(ComercialLeadChannel).where(ComercialLeadChannel.lead_id == lead_id, ComercialLeadChannel.tipo == tipo))).scalar_one_or_none()
    if x:
        x.valor = valor; x.ativo = True
    else:
        x = ComercialLeadChannel(lead_id=lead_id, tipo=tipo, valor=valor, ativo=True); db.add(x)
    await activity(db, lead_id, user.id, "canal_atualizado", f"Canal atualizado: {tipo}")
    await db.commit(); await db.refresh(x)
    return d(x)


@router.delete("/leads/{lead_id}/canais/{canal_id}", status_code=204)
async def excluir_canal(lead_id: int, canal_id: int, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialLeadChannel).where(ComercialLeadChannel.id == canal_id, ComercialLeadChannel.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Canal não encontrado")
    x.ativo = False; await activity(db, lead_id, user.id, "canal_removido", f"Canal removido: {x.tipo}"); await db.commit()


@router.post("/leads/{lead_id}/fontes", status_code=201)
async def criar_fonte(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    evidencia = (payload.get("evidencia") or "").strip()
    if not evidencia: raise HTTPException(400, "Trecho/evidência é obrigatório")
    x = ComercialAuditSource(
        lead_id=lead_id,
        tipo=payload.get("tipo", "outro"),
        fonte=payload.get("fonte"),
        url=payload.get("url"),
        evidencia=evidencia,
        consultado_em=_parse_datetime(payload.get("consultado_em")) or _now(),
    )
    db.add(x); await activity(db, lead_id, user.id, "fonte_adicionada", f"Fonte adicionada: {x.tipo}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.delete("/leads/{lead_id}/fontes/{fonte_id}", status_code=204)
async def excluir_fonte(lead_id: int, fonte_id: int, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialAuditSource).where(ComercialAuditSource.id == fonte_id, ComercialAuditSource.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Fonte não encontrada")
    await db.delete(x); await activity(db, lead_id, user.id, "fonte_removida", "Fonte/evidência removida"); await db.commit()


@router.post("/leads/{lead_id}/ouros", status_code=201)
async def criar_ouro(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    if not (payload.get("titulo") or "").strip() or not (payload.get("evidencia") or "").strip():
        raise HTTPException(400, "Título e evidência são obrigatórios")
    x = ComercialInsight(
        lead_id=lead_id,
        categoria=payload.get("categoria", "outro"),
        titulo=payload.get("titulo").strip(),
        evidencia=payload.get("evidencia").strip(),
        impacto=(payload.get("impacto") or "").strip(),
        fonte=payload.get("fonte"),
        grau_confianca=payload.get("grau_confianca", "medio"),
        utilizado=bool(payload.get("utilizado", False)),
    )
    db.add(x); await activity(db, lead_id, user.id, "ouro_adicionado", f"Ouro adicionado: {x.titulo}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.patch("/leads/{lead_id}/ouros/{ouro_id}")
async def editar_ouro(lead_id: int, ouro_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialInsight).where(ComercialInsight.id == ouro_id, ComercialInsight.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Ouro não encontrado")
    for k in ("categoria", "titulo", "evidencia", "impacto", "fonte", "grau_confianca", "utilizado", "ativo"):
        if k in payload: setattr(x, k, payload[k])
    if "utilizado" in payload:
        x.utilizado_em = _now() if payload["utilizado"] else None
    await activity(db, lead_id, user.id, "ouro_atualizado", f"Ouro atualizado: {x.titulo}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.post("/leads/{lead_id}/memoria", status_code=201)
async def criar_memoria(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db, lead_id)
    conteudo = (payload.get("conteudo") or "").strip()
    if not conteudo: raise HTTPException(400, "Conteúdo é obrigatório")
    tipo = payload.get("tipo", "fato")
    if tipo not in TIPOS_FATO: raise HTTPException(400, "Tipo de memória inválido")
    x = ComercialLeadFact(lead_id=lead_id, tipo=tipo, conteudo=conteudo, fonte_id=payload.get("fonte_id"), ativo=True)
    db.add(x); await activity(db, lead_id, user.id, "memoria_adicionada", f"Memória adicionada: {tipo}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.delete("/leads/{lead_id}/memoria/{fact_id}", status_code=204)
async def excluir_memoria(lead_id: int, fact_id: int, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialLeadFact).where(ComercialLeadFact.id == fact_id, ComercialLeadFact.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Memória não encontrada")
    x.ativo = False; await activity(db, lead_id, user.id, "memoria_removida", "Item de memória removido"); await db.commit()


@router.post("/leads/{lead_id}/analisar")
async def analisar(lead_id: int, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    old_status = lead.status_codigo
    lead.status_codigo = "em_analise"
    await db.flush()
    # updated_at é calculado pelo banco no UPDATE; após o flush o SQLAlchemy pode
    # expirar esse atributo. Recarregar evita lazy-load síncrono/MissingGreenlet
    # quando o lead inteiro é serializado para o contexto da IA.
    await db.refresh(lead)
    ctx = await contexto(db, lead)
    if not ctx["fontes"] and not ctx["canais"]:
        lead.status_codigo = "a_pesquisar"
        await db.commit()
        raise HTTPException(409, "Adicione ao menos um canal ou uma evidência pública antes de analisar. A IA não deve inventar uma auditoria sem dados.")
    formato = '{"resumo":"","tese_central":"","forcas":[],"vazamentos":[],"oportunidades":[],"ouros":[{"categoria":"google|instagram|cardapio|delivery|reputacao|posicionamento|conversao|reserva|whatsapp|seo|site|prova_social|recorrencia|ticket_medio|experiencia|outro","titulo":"","evidencia":"","impacto":"","fonte":"","grau_confianca":"alto|medio|baixo"}],"prioridade":"","risco_comercial":"","observacoes":[],"fit_score":0,"opportunity_score":0}'
    try:
        result, usage = await gerar_json(
            "Audite o negócio usando somente o contexto fornecido. Não trate ausência de dado como problema observado. Separe fatos de hipóteses e gere ouros utilizáveis em conversa comercial progressiva.",
            ctx,
            formato,
        )
    except RuntimeError as exc:
        lead.status_codigo = old_status if old_status != "em_analise" else "a_pesquisar"
        await db.commit()
        raise HTTPException(503, str(exc))
    audit = ComercialAudit(
        lead_id=lead.id,
        executado_por_envoxer_id=user.id,
        resumo=result.get("resumo"),
        tese_central=result.get("tese_central"),
        forcas=result.get("forcas", []),
        vazamentos=result.get("vazamentos", []),
        oportunidades=result.get("oportunidades", []),
        prioridade=result.get("prioridade"),
        risco_comercial=result.get("risco_comercial"),
        observacoes=result.get("observacoes", []),
        modelo=settings.OPENAI_MODEL,
    )
    db.add(audit); await db.flush()
    for ouro in result.get("ouros", []):
        if not ouro.get("evidencia"):
            continue
        db.add(ComercialInsight(
            lead_id=lead.id,
            audit_id=audit.id,
            categoria=ouro.get("categoria", "outro"),
            titulo=ouro.get("titulo") or "Ouro",
            evidencia=ouro.get("evidencia", ""),
            impacto=ouro.get("impacto", ""),
            fonte=ouro.get("fonte"),
            grau_confianca=ouro.get("grau_confianca", "medio"),
        ))
    lead.fit_score = max(0, min(100, int(result.get("fit_score", 0) or 0)))
    lead.opportunity_score = max(0, min(100, int(result.get("opportunity_score", 0) or 0)))
    lead.score_total = round((lead.fit_score + lead.opportunity_score + lead.engagement_score) / 3)
    lead.status_codigo = "pronto_abordagem"
    lead.proxima_acao = "Gerar primeira abordagem"
    lead.proxima_acao_em = _now()
    db.add(ComercialAIGeneration(
        lead_id=lead.id,
        usuario_envoxer_id=user.id,
        tipo="auditoria",
        modelo=settings.OPENAI_MODEL,
        prompt_version=PROMPT_VERSION,
        entrada=ctx,
        resultado=result,
        tokens_entrada=usage.get("input_tokens"),
        tokens_saida=usage.get("output_tokens"),
    ))
    await activity(db, lead.id, user.id, "auditoria_ia", "Auditoria por IA concluída")
    await db.commit()
    return result


@router.post("/leads/{lead_id}/gerar-mensagem")
async def gerar_mensagem(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    if lead.nao_prospectar:
        raise HTTPException(409, "Este lead está marcado como NÃO PROSPECTAR")
    canal = payload.get("canal", "instagram")
    ouro_id = payload.get("insight_id")
    ouro = None
    if ouro_id:
        ouro = (await db.execute(select(ComercialInsight).where(
            ComercialInsight.id == int(ouro_id), ComercialInsight.lead_id == lead_id, ComercialInsight.ativo.is_(True)
        ))).scalar_one_or_none()
    if ouro is None:
        ouro = (await db.execute(select(ComercialInsight).where(
            ComercialInsight.lead_id == lead_id,
            ComercialInsight.utilizado.is_(False),
            ComercialInsight.ativo.is_(True),
        ).order_by(ComercialInsight.created_at))).scalars().first()
    if not ouro:
        raise HTTPException(409, "Não há ouro disponível. Adicione evidências/ouros ou rode uma nova auditoria.")
    ctx = await contexto(db, lead)
    ctx["ouro_escolhido"] = d(ouro)
    ctx["canal"] = canal
    result, usage = await gerar_json(
        "Gere somente a próxima mensagem. Se for primeira abordagem, o objetivo é gerar resposta, usando apenas um ouro. Não peça reunião na primeira mensagem.",
        ctx,
        '{"mensagem":"","gancho":"","evidencia":"","pepita":"","marcacao_valor":"","pergunta":""}',
    )
    db.add(ComercialAIGeneration(
        lead_id=lead.id,
        usuario_envoxer_id=user.id,
        tipo="mensagem",
        modelo=settings.OPENAI_MODEL,
        prompt_version=PROMPT_VERSION,
        entrada=ctx,
        resultado=result,
        tokens_entrada=usage.get("input_tokens"),
        tokens_saida=usage.get("output_tokens"),
    ))
    await activity(db, lead.id, user.id, "mensagem_gerada", f"Mensagem gerada com ouro reservado: {ouro.titulo}", {"insight_id": ouro.id, "canal": canal})
    await db.commit()
    result["insight_id"] = ouro.id
    result["canal"] = canal
    return result


@router.post("/leads/{lead_id}/mensagens", status_code=201)
async def registrar_mensagem(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    if lead.nao_prospectar and payload.get("direcao", "enviada") == "enviada":
        raise HTTPException(409, "Este lead está na lista de não prospectar. O envio não pode ser registrado como nova prospecção.")
    conteudo = (payload.get("conteudo") or "").strip()
    if not conteudo:
        raise HTTPException(400, "Conteúdo da mensagem é obrigatório")
    direcao = payload.get("direcao", "enviada")
    before = lead.status_codigo
    now = _now()
    insight_id = payload.get("insight_id")
    if direcao == "recebida":
        lead.status_codigo = "respondeu"
        lead.engagement_score = min(100, lead.engagement_score + 25)
        lead.temperatura = "quente" if lead.engagement_score >= 60 else "morno"
        lead.proxima_acao = "Responder lead"
        lead.proxima_acao_em = now
        await _close_or_pause_cadence(db, lead_id, "pausada")
        obj = payload.get("objecao_categoria")
        if obj:
            db.add(ComercialLeadFact(lead_id=lead_id, tipo="objecao", conteudo=f"Objeção registrada: {obj}", ativo=True))
    else:
        current_step = await _advance_active_cadence(db, lead)
        if current_step is None and before in {"novo", "a_pesquisar", "em_analise", "analise_concluida", "pronto_abordagem"}:
            lead.status_codigo = "mensagem_1_enviada"
            lead.proxima_acao = "Follow-up com nova pepita"
            lead.proxima_acao_em = now + timedelta(days=2)
        if insight_id:
            ouro = (await db.execute(select(ComercialInsight).where(ComercialInsight.id == int(insight_id), ComercialInsight.lead_id == lead_id))).scalar_one_or_none()
            if ouro:
                ouro.utilizado = True
                ouro.utilizado_em = now
    lead.ultimo_contato_em = now
    lead.score_total = round((lead.fit_score + lead.opportunity_score + lead.engagement_score) / 3)
    msg = ComercialMessage(
        lead_id=lead_id,
        direcao=direcao,
        canal=payload.get("canal", "instagram"),
        tipo=payload.get("tipo", "abordagem" if direcao == "enviada" else "resposta"),
        conteudo=conteudo,
        responsavel_envoxer_id=user.id,
        insight_id=insight_id,
        objecao_categoria=payload.get("objecao_categoria"),
        status_antes=before,
        status_depois=lead.status_codigo,
        enviado_em=_parse_datetime(payload.get("enviado_em")) or now,
        status_envio=payload.get("status_envio", "registrada"),
        gerado_por_ia=bool(payload.get("gerado_por_ia", False)),
        aprovado_em=now if payload.get("gerado_por_ia") and direcao == "enviada" else None,
        metadados=payload.get("metadados"),
    )
    db.add(msg)
    await activity(db, lead_id, user.id, "mensagem_" + direcao, f"Mensagem {direcao} registrada", {"canal": msg.canal, "tipo": msg.tipo, "insight_id": insight_id})
    await db.commit(); await db.refresh(msg)
    return d(msg)


@router.post("/leads/{lead_id}/gerar-resposta")
async def gerar_resposta(lead_id: int, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    ctx = await contexto(db, lead)
    received = [m for m in ctx["mensagens"] if m["direcao"] == "recebida"]
    if not received:
        raise HTTPException(409, "Registre a resposta do lead antes de gerar a próxima resposta.")
    result, usage = await gerar_json(
        "Interprete a resposta mais recente e gere SOMENTE a próxima mensagem. Se houver um novo ouro disponível, use no máximo um e apenas se fizer sentido. Nunca discuta com o lead.",
        ctx,
        '{"interpretacao":{"tipo":"duvida|abertura|objecao|negativa|outro","sentimento":"positivo|neutro|negativo","intencao":"","objecao":"","proximo_passo":""},"mensagem":"","insight_id":null}',
    )
    iid = result.get("insight_id")
    if iid:
        valid = (await db.execute(select(ComercialInsight.id).where(
            ComercialInsight.id == int(iid), ComercialInsight.lead_id == lead_id, ComercialInsight.utilizado.is_(False), ComercialInsight.ativo.is_(True)
        ))).scalar_one_or_none()
        if not valid:
            result["insight_id"] = None
    db.add(ComercialAIGeneration(
        lead_id=lead.id,
        usuario_envoxer_id=user.id,
        tipo="proxima_resposta",
        modelo=settings.OPENAI_MODEL,
        prompt_version=PROMPT_VERSION,
        entrada=ctx,
        resultado=result,
        tokens_entrada=usage.get("input_tokens"),
        tokens_saida=usage.get("output_tokens"),
    ))
    await activity(db, lead.id, user.id, "resposta_gerada", "Próxima resposta gerada pela IA")
    await db.commit()
    return result


@router.post("/leads/{lead_id}/tarefas", status_code=201)
async def criar_tarefa(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    titulo = (payload.get("titulo") or "").strip()
    if not titulo: raise HTTPException(400, "Título da tarefa é obrigatório")
    x = ComercialTask(
        lead_id=lead_id,
        titulo=titulo,
        descricao=payload.get("descricao"),
        responsavel_envoxer_id=payload.get("responsavel_envoxer_id") or lead.responsavel_envoxer_id or user.id,
        prazo=_parse_datetime(payload.get("prazo")),
        status="pendente",
        tipo=payload.get("tipo", "comercial"),
    )
    db.add(x); await activity(db, lead_id, user.id, "tarefa_criada", f"Tarefa criada: {titulo}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.patch("/tarefas/{task_id}")
async def editar_tarefa(task_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialTask).where(ComercialTask.id == task_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Tarefa não encontrada")
    for k in ("titulo", "descricao", "responsavel_envoxer_id", "status", "tipo"):
        if k in payload: setattr(x, k, payload[k])
    if "prazo" in payload: x.prazo = _parse_datetime(payload.get("prazo"))
    if "status" in payload:
        x.concluida_em = _now() if payload["status"] == "concluida" else None
    await activity(db, x.lead_id, user.id, "tarefa_atualizada", f"Tarefa atualizada: {x.titulo}", {"status": x.status})
    await db.commit(); await db.refresh(x)
    return d(x)


@router.get("/tarefas")
async def listar_tarefas(status: Optional[str] = None, vencidas: Optional[bool] = None, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    q = select(ComercialTask, ComercialLead).join(ComercialLead, ComercialLead.id == ComercialTask.lead_id).where(ComercialLead.deleted_at.is_(None))
    if status: q = q.where(ComercialTask.status == status)
    if vencidas: q = q.where(ComercialTask.status == "pendente", ComercialTask.prazo < _now())
    if user.permissao == "envoxer": q = q.where(or_(ComercialTask.responsavel_envoxer_id == user.id, ComercialTask.responsavel_envoxer_id.is_(None)))
    rows = (await db.execute(q.order_by(ComercialTask.status, ComercialTask.prazo.asc().nullslast()))).all()
    resp = await _responsaveis(db)
    return [{**d(t), "lead_nome": l.nome_estabelecimento, "lead_status": l.status_codigo, "responsavel_nome": resp.get(t.responsavel_envoxer_id).nome if resp.get(t.responsavel_envoxer_id) else None} for t, l in rows]


@router.get("/cadencias")
async def listar_cadencias(db: AsyncSession = Depends(get_db), _: Envoxer = Depends(get_current_envoxer)):
    rows = (await db.execute(select(ComercialCadence).order_by(ComercialCadence.ativa.desc(), ComercialCadence.nome))).scalars().all()
    out = []
    for c in rows:
        steps = [d(x) for x in (await db.execute(select(ComercialCadenceStep).where(ComercialCadenceStep.cadence_id == c.id).order_by(ComercialCadenceStep.ordem))).scalars().all()]
        ativos = (await db.execute(select(func.count()).select_from(ComercialLeadCadence).where(ComercialLeadCadence.cadence_id == c.id, ComercialLeadCadence.status == "ativa"))).scalar_one()
        out.append({**d(c), "steps": steps, "leads_ativos": ativos})
    return out


@router.post("/cadencias", status_code=201)
async def criar_cadencia(payload: dict, db: AsyncSession = Depends(get_db), _: Envoxer = Depends(get_current_gestor_ou_admin)):
    nome = (payload.get("nome") or "").strip()
    if not nome: raise HTTPException(400, "Nome da cadência é obrigatório")
    steps = payload.get("steps") or []
    if not steps: raise HTTPException(400, "Cadência precisa ter ao menos uma etapa")
    c = ComercialCadence(nome=nome, descricao=payload.get("descricao"), ativa=bool(payload.get("ativa", True)))
    db.add(c); await db.flush()
    for i, step in enumerate(steps, 1):
        db.add(ComercialCadenceStep(cadence_id=c.id, ordem=int(step.get("ordem", i)), dias_apos_inicio=int(step.get("dias_apos_inicio", 0)), tipo=step.get("tipo", "follow_up"), orientacao=step.get("orientacao")))
    await db.commit(); await db.refresh(c)
    return d(c)


@router.patch("/cadencias/{cadence_id}")
async def editar_cadencia(cadence_id: int, payload: dict, db: AsyncSession = Depends(get_db), _: Envoxer = Depends(get_current_gestor_ou_admin)):
    c = (await db.execute(select(ComercialCadence).where(ComercialCadence.id == cadence_id))).scalar_one_or_none()
    if not c: raise HTTPException(404, "Cadência não encontrada")
    for k in ("nome", "descricao", "ativa"):
        if k in payload: setattr(c, k, payload[k])
    if "steps" in payload:
        await db.execute(delete(ComercialCadenceStep).where(ComercialCadenceStep.cadence_id == cadence_id))
        for i, step in enumerate(payload.get("steps") or [], 1):
            db.add(ComercialCadenceStep(cadence_id=c.id, ordem=int(step.get("ordem", i)), dias_apos_inicio=int(step.get("dias_apos_inicio", 0)), tipo=step.get("tipo", "follow_up"), orientacao=step.get("orientacao")))
    await db.commit(); await db.refresh(c)
    return d(c)


@router.post("/leads/{lead_id}/cadencia")
async def iniciar_cadencia(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    if lead.nao_prospectar: raise HTTPException(409, "Lead está na lista de não prospectar")
    active = (await db.execute(select(ComercialLeadCadence).where(ComercialLeadCadence.lead_id == lead_id, ComercialLeadCadence.status == "ativa"))).scalars().first()
    if active: raise HTTPException(409, "Este lead já possui uma cadência ativa")
    cadence_id = int(payload.get("cadence_id", 1))
    cadence = (await db.execute(select(ComercialCadence).where(ComercialCadence.id == cadence_id, ComercialCadence.ativa.is_(True)))).scalar_one_or_none()
    if not cadence: raise HTTPException(404, "Cadência não encontrada ou inativa")
    first = (await db.execute(select(ComercialCadenceStep).where(ComercialCadenceStep.cadence_id == cadence_id).order_by(ComercialCadenceStep.ordem))).scalars().first()
    if not first: raise HTTPException(409, "Cadência sem etapas")
    start = _now()
    due = start + timedelta(days=first.dias_apos_inicio)
    lc = ComercialLeadCadence(lead_id=lead_id, cadence_id=cadence_id, status="ativa", iniciada_em=start, proximo_step_ordem=first.ordem, proximo_step_em=due)
    db.add(lc); lead.proxima_acao="Enviar mensagem inicial" if first.tipo == "abordagem" else (first.orientacao or "Executar etapa da cadência"); lead.proxima_acao_em=due
    await activity(db, lead_id, user.id, "cadencia_iniciada", f"Cadência iniciada: {cadence.nome}"); await db.commit(); await db.refresh(lc)
    return d(lc)


@router.patch("/leads/{lead_id}/cadencia/{lead_cadence_id}")
async def editar_lead_cadencia(lead_id: int, lead_cadence_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialLeadCadence).where(ComercialLeadCadence.id == lead_cadence_id, ComercialLeadCadence.lead_id == lead_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Cadência do lead não encontrada")
    action = payload.get("status")
    if action not in {"ativa", "pausada", "encerrada"}: raise HTTPException(400, "Status inválido")
    x.status = action
    if action == "pausada": x.pausada_em = _now()
    if action == "encerrada": x.encerrada_em = _now(); x.proximo_step_em = None
    if action == "ativa": x.pausada_em = None
    await activity(db, lead_id, user.id, "cadencia_atualizada", f"Cadência: {action}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.post("/leads/{lead_id}/oportunidades", status_code=201)
async def criar_oportunidade(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    x = ComercialOpportunity(
        lead_id=lead_id,
        contato=payload.get("contato"),
        dor_principal=payload.get("dor_principal"),
        solucao_provavel=payload.get("solucao_provavel"),
        servicos_relacionados=payload.get("servicos_relacionados") or [],
        valor_estimado=payload.get("valor_estimado") or None,
        probabilidade=max(0, min(100, int(payload.get("probabilidade", 0) or 0))),
        etapa=payload.get("etapa", "oportunidade"),
        proxima_acao=payload.get("proxima_acao"),
        data_prevista=_parse_date(payload.get("data_prevista")),
        responsavel_envoxer_id=payload.get("responsavel_envoxer_id") or lead.responsavel_envoxer_id or user.id,
    )
    db.add(x)
    old = lead.status_codigo
    lead.status_codigo = "oportunidade"
    if x.proxima_acao:
        lead.proxima_acao = x.proxima_acao
    await _close_or_pause_cadence(db, lead_id, "encerrada")
    await activity(db, lead_id, user.id, "oportunidade_criada", "Oportunidade comercial criada", {"antes": old, "valor_estimado": str(x.valor_estimado) if x.valor_estimado else None})
    await db.commit(); await db.refresh(x)
    return d(x)


@router.patch("/oportunidades/{opportunity_id}")
async def editar_oportunidade(opportunity_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    x = (await db.execute(select(ComercialOpportunity).where(ComercialOpportunity.id == opportunity_id))).scalar_one_or_none()
    if not x: raise HTTPException(404, "Oportunidade não encontrada")
    for k in ("contato", "dor_principal", "solucao_provavel", "servicos_relacionados", "probabilidade", "etapa", "proxima_acao", "responsavel_envoxer_id"):
        if k in payload: setattr(x, k, payload[k])
    if "valor_estimado" in payload: x.valor_estimado = payload.get("valor_estimado") or None
    if "data_prevista" in payload: x.data_prevista = _parse_date(payload.get("data_prevista"))
    lead = await lead_or_404(db, x.lead_id)
    if x.etapa in {s for s in STATUS_REUNIAO_OU_DEPOIS | {"oportunidade", "perdido"}}:
        old = lead.status_codigo; lead.status_codigo = x.etapa; await _status_side_effects(db, lead, old, lead.status_codigo)
    if x.proxima_acao: lead.proxima_acao = x.proxima_acao
    await activity(db, x.lead_id, user.id, "oportunidade_atualizada", f"Oportunidade atualizada: {x.etapa}"); await db.commit(); await db.refresh(x)
    return d(x)


@router.get("/oportunidades")
async def listar_oportunidades(db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    q = select(ComercialOpportunity, ComercialLead).join(ComercialLead, ComercialLead.id == ComercialOpportunity.lead_id).where(ComercialLead.deleted_at.is_(None))
    if user.permissao == "envoxer": q = q.where(or_(ComercialOpportunity.responsavel_envoxer_id == user.id, ComercialOpportunity.responsavel_envoxer_id.is_(None)))
    rows = (await db.execute(q.order_by(ComercialOpportunity.updated_at.desc()))).all(); resp=await _responsaveis(db)
    return [{**d(o), "lead_nome": l.nome_estabelecimento, "lead_status": l.status_codigo, "responsavel_nome": resp.get(o.responsavel_envoxer_id).nome if resp.get(o.responsavel_envoxer_id) else None} for o,l in rows]


@router.get("/conversas")
async def conversas(db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    q = select(ComercialLead).where(ComercialLead.deleted_at.is_(None))
    if user.permissao == "envoxer": q = q.where(or_(ComercialLead.responsavel_envoxer_id == user.id, ComercialLead.responsavel_envoxer_id.is_(None)))
    leads = (await db.execute(q)).scalars().all(); ids=[x.id for x in leads]
    if not ids: return []
    msgs=(await db.execute(select(ComercialMessage).where(ComercialMessage.lead_id.in_(ids)).order_by(ComercialMessage.enviado_em.desc()))).scalars().all()
    grouped=defaultdict(list)
    for m in msgs: grouped[m.lead_id].append(m)
    resp=await _responsaveis(db); tags=await _tags_por_lead(db,ids); out=[]
    for lead in leads:
        ms=grouped.get(lead.id,[])
        if not ms: continue
        last=ms[0]
        unread_action=lead.status_codigo=="respondeu"
        out.append({"lead":await _lead_view(db,lead,resp,tags),"ultima_mensagem":d(last),"total":len(ms),"aguardando_acao":unread_action})
    out.sort(key=lambda x:x["ultima_mensagem"]["enviado_em"], reverse=True)
    return out


@router.post("/leads/{lead_id}/nao-prospectar")
async def nao_prospectar(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    lead = await lead_or_404(db, lead_id)
    motivo = payload.get("motivo") or "Pedido para não contatar"
    lead.nao_prospectar=True; lead.status_codigo="nao_prospectar"; lead.proxima_acao=None; lead.proxima_acao_em=None
    await _suppression_add(db,lead,user.id,motivo); await _close_or_pause_cadence(db,lead_id,"encerrada"); await activity(db,lead_id,user.id,"nao_prospectar","Lead marcado como NÃO PROSPECTAR",{"motivo":motivo}); await db.commit()
    return {"ok":True}


@router.post("/tags")
async def criar_tag(payload: dict, db: AsyncSession = Depends(get_db), _: Envoxer = Depends(get_current_envoxer)):
    nome=(payload.get("nome") or "").strip()
    if not nome: raise HTTPException(400,"Nome da tag é obrigatório")
    x=(await db.execute(select(ComercialTag).where(func.lower(ComercialTag.nome)==nome.lower()))).scalar_one_or_none()
    if not x: x=ComercialTag(nome=nome); db.add(x); await db.commit(); await db.refresh(x)
    return d(x)


async def _set_tag(db: AsyncSession, lead_id: int, nome: str):
    nome=nome.strip()
    if not nome: return
    tag=(await db.execute(select(ComercialTag).where(func.lower(ComercialTag.nome)==nome.lower()))).scalar_one_or_none()
    if not tag: tag=ComercialTag(nome=nome); db.add(tag); await db.flush()
    exists=(await db.execute(select(ComercialLeadTag).where(ComercialLeadTag.lead_id==lead_id,ComercialLeadTag.tag_id==tag.id))).scalar_one_or_none()
    if not exists: db.add(ComercialLeadTag(lead_id=lead_id,tag_id=tag.id))


@router.post("/leads/{lead_id}/tags")
async def adicionar_tag(lead_id: int, payload: dict, db: AsyncSession = Depends(get_db), user: Envoxer = Depends(get_current_envoxer)):
    await lead_or_404(db,lead_id); nome=(payload.get("nome") or "").strip()
    if not nome: raise HTTPException(400,"Nome da tag é obrigatório")
    await _set_tag(db,lead_id,nome); await activity(db,lead_id,user.id,"tag_adicionada",f"Tag adicionada: {nome}"); await db.commit(); return {"ok":True}


@router.delete("/leads/{lead_id}/tags/{tag_id}",status_code=204)
async def remover_tag(lead_id:int,tag_id:int,db:AsyncSession=Depends(get_db),user:Envoxer=Depends(get_current_envoxer)):
    await db.execute(delete(ComercialLeadTag).where(ComercialLeadTag.lead_id==lead_id,ComercialLeadTag.tag_id==tag_id)); await activity(db,lead_id,user.id,"tag_removida","Tag removida"); await db.commit()


def _read_import(raw: bytes, name: str):
    name=name.lower()
    if name.endswith(".csv"):
        text=raw.decode("utf-8-sig",errors="replace")
        sample=text[:5000]
        try: dialect=csv.Sniffer().sniff(sample,delimiters=",;\t")
        except csv.Error: dialect=csv.excel
        rows=list(csv.DictReader(io.StringIO(text),dialect=dialect))
        return [{str(k): (v if v is not None else "") for k,v in row.items()} for row in rows]
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(500,"Leitor XLSX não instalado no backend")
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        ws=wb.active
        values=list(ws.iter_rows(values_only=True))
        if not values: return []
        headers=[str(v or "").strip() for v in values[0]]
        out=[]
        for line in values[1:]:
            if not any(v not in (None,"") for v in line): continue
            out.append({headers[i] if i<len(headers) else f"col_{i+1}": ("" if v is None else v.isoformat() if isinstance(v,(date,datetime)) else str(v)) for i,v in enumerate(line)})
        return out
    raise HTTPException(400,"Envie um arquivo CSV ou XLSX")


IMPORT_ALIASES={
    "data":"data_entrada","data de entrada":"data_entrada","nome":"nome_estabelecimento","estabelecimento":"nome_estabelecimento","nome do estabelecimento":"nome_estabelecimento","nome fantasia":"nome_fantasia","segmento":"segmento","subsegmento":"subsegmento","cidade":"cidade","estado":"estado","bairro":"bairro","instagram":"instagram","facebook":"facebook","tiktok":"tiktok","site":"site","google maps":"google_maps","tripadvisor":"tripadvisor","ifood":"ifood","rappi":"rappi","99food":"99food","linktree":"linktree","cardapio":"cardapio_digital","cardápio":"cardapio_digital","cardápio digital":"cardapio_digital","telefone":"telefone","whatsapp":"whatsapp","e-mail":"email","email":"email","contato":"contato_nome","nome do contato":"contato_nome","cargo":"contato_cargo","origem":"origem","observacoes":"observacoes","observações":"observacoes"
}
IMPORT_FIELDS=["data_entrada","nome_estabelecimento","nome_fantasia","segmento","subsegmento","cidade","estado","bairro","instagram","facebook","tiktok","site","google_maps","tripadvisor","ifood","rappi","99food","linktree","cardapio_digital","telefone","whatsapp","email","contato_nome","contato_cargo","contato_telefone","contato_whatsapp","contato_email","origem","observacoes"]


async def _find_duplicate(db:AsyncSession,row:dict,mapping:dict):
    def val(field):
        col=next((c for c,f in mapping.items() if f==field),None); return str(row.get(col,"" ) or "").strip() if col else ""
    email=val("email"); whatsapp=_phone(val("whatsapp")); telefone=_phone(val("telefone")); nome=val("nome_estabelecimento"); cidade=val("cidade"); instagram=val("instagram")
    if email:
        x=(await db.execute(select(ComercialLead).where(ComercialLead.deleted_at.is_(None),func.lower(ComercialLead.email)==email.lower()))).scalars().first()
        if x:return x,"email"
    for raw,field in ((whatsapp,"whatsapp"),(telefone,"telefone")):
        if raw:
            leads=(await db.execute(select(ComercialLead).where(ComercialLead.deleted_at.is_(None),getattr(ComercialLead,field).is_not(None)))).scalars().all()
            x=next((l for l in leads if _phone(getattr(l,field))==raw),None)
            if x:return x,field
    if instagram:
        ch=(await db.execute(select(ComercialLeadChannel).where(ComercialLeadChannel.tipo=="instagram",ComercialLeadChannel.ativo.is_(True),func.lower(ComercialLeadChannel.valor)==instagram.lower()))).scalars().first()
        if ch:
            x=await lead_or_404(db,ch.lead_id); return x,"instagram"
    if nome:
        q=select(ComercialLead).where(ComercialLead.deleted_at.is_(None),func.lower(ComercialLead.nome_estabelecimento)==nome.lower())
        if cidade:q=q.where(func.lower(ComercialLead.cidade)==cidade.lower())
        x=(await db.execute(q)).scalars().first()
        if x:return x,"nome+cidade" if cidade else "nome"
    return None,None


@router.post("/importar/preview")
async def importar_preview(arquivo:UploadFile=File(...),db:AsyncSession=Depends(get_db),_:Envoxer=Depends(get_current_envoxer)):
    raw=await arquivo.read(); rows=_read_import(raw,arquivo.filename or "")
    if len(rows)>5000: raise HTTPException(400,"Limite do importador: 5.000 linhas por arquivo")
    cols=list(rows[0].keys()) if rows else []
    mapping={c:IMPORT_ALIASES.get(_norm(c)) for c in cols}
    duplicates=[]
    for idx,row in enumerate(rows[:200]):
        lead,reason=await _find_duplicate(db,row,mapping)
        if lead:duplicates.append({"linha":idx+2,"lead_id":lead.id,"lead_nome":lead.nome_estabelecimento,"motivo":reason})
    return {"colunas":cols,"mapeamento_sugerido":mapping,"campos_destino":IMPORT_FIELDS,"preview":rows[:20],"rows":rows,"total":len(rows),"duplicados":duplicates,"duplicados_aviso":"A prévia verifica até 200 linhas; a importação valida todas novamente."}


@router.post("/importar/commit")
async def importar_commit(payload:dict,db:AsyncSession=Depends(get_db),user:Envoxer=Depends(get_current_envoxer)):
    rows=payload.get("rows") or []; mapping=payload.get("mapping") or {}; strategy=payload.get("duplicate_strategy","skip")
    if not rows or not mapping: raise HTTPException(400,"Nenhuma linha ou mapeamento recebido")
    if len(rows)>5000: raise HTTPException(400,"Limite de 5.000 linhas")
    if "nome_estabelecimento" not in mapping.values(): raise HTTPException(400,"Mapeie a coluna NOME para Nome do estabelecimento")
    created=updated=skipped=0; errors=[]
    for idx,row in enumerate(rows):
        values={dest:str(row.get(src,"") or "").strip() for src,dest in mapping.items() if dest and src in row}
        if not values.get("nome_estabelecimento"):
            skipped+=1; errors.append({"linha":idx+2,"erro":"nome vazio"}); continue
        duplicate,reason=await _find_duplicate(db,row,mapping)
        lead=duplicate
        if duplicate and strategy=="skip": skipped+=1; continue
        lead_fields={k:v for k,v in values.items() if k in {"nome_estabelecimento","nome_fantasia","segmento","subsegmento","cidade","estado","bairro","telefone","whatsapp","email","origem","observacoes"} and v!=""}
        if "data_entrada" in values: lead_fields["data_entrada"]=_parse_date(values["data_entrada"]) or date.today()
        if not lead_fields.get("origem"): lead_fields["origem"]="importacao"
        if duplicate and strategy=="update":
            for k,v in lead_fields.items(): setattr(lead,k,v)
            updated+=1; await activity(db,lead.id,user.id,"lead_importado_atualizado",f"Lead atualizado por importação ({reason})")
        else:
            lead_fields.setdefault("proxima_acao", "Pesquisar e analisar lead")
            lead_fields.setdefault("proxima_acao_em", _now())
            lead=ComercialLead(**lead_fields,responsavel_envoxer_id=user.id if user.permissao=="envoxer" else None); db.add(lead); await db.flush(); created+=1; await activity(db,lead.id,user.id,"lead_importado","Lead criado por importação")
        for tipo in CANAIS:
            if values.get(tipo):
                ch=(await db.execute(select(ComercialLeadChannel).where(ComercialLeadChannel.lead_id==lead.id,ComercialLeadChannel.tipo==tipo))).scalar_one_or_none()
                if ch: ch.valor=values[tipo]; ch.ativo=True
                else: db.add(ComercialLeadChannel(lead_id=lead.id,tipo=tipo,valor=values[tipo],ativo=True))
        if values.get("contato_nome"):
            existing=(await db.execute(select(ComercialLeadContact).where(ComercialLeadContact.lead_id==lead.id,func.lower(ComercialLeadContact.nome)==values["contato_nome"].lower()))).scalar_one_or_none()
            cdata={"nome":values["contato_nome"],"cargo":values.get("contato_cargo") or None,"telefone":values.get("contato_telefone") or None,"whatsapp":values.get("contato_whatsapp") or None,"email":values.get("contato_email") or None}
            if existing:
                for k,v in cdata.items(): setattr(existing,k,v)
            else: db.add(ComercialLeadContact(lead_id=lead.id,principal=True,**cdata))
    await db.commit()
    return {"criados":created,"atualizados":updated,"ignorados":skipped,"erros":errors[:50],"total":len(rows)}


@router.post("/leads/bulk")
async def bulk_leads(payload:dict,db:AsyncSession=Depends(get_db),user:Envoxer=Depends(get_current_envoxer)):
    ids=[int(x) for x in payload.get("ids") or []]
    if not ids: raise HTTPException(400,"Selecione ao menos um lead")
    leads=(await db.execute(select(ComercialLead).where(ComercialLead.id.in_(ids),ComercialLead.deleted_at.is_(None)))).scalars().all()
    action=payload.get("action"); value=payload.get("value")
    for lead in leads:
        if action=="status":
            old=lead.status_codigo; lead.status_codigo=str(value); await _status_side_effects(db,lead,old,lead.status_codigo); await activity(db,lead.id,user.id,"status_alterado",f"Status: {old} → {lead.status_codigo}")
        elif action=="responsavel": lead.responsavel_envoxer_id=int(value) if value else None; await activity(db,lead.id,user.id,"responsavel_alterado","Responsável comercial alterado em massa")
        elif action=="nao_prospectar": lead.nao_prospectar=True; lead.status_codigo="nao_prospectar"; lead.proxima_acao=None; lead.proxima_acao_em=None; await _suppression_add(db,lead,user.id,"Ação em massa"); await _close_or_pause_cadence(db,lead.id,"encerrada"); await activity(db,lead.id,user.id,"nao_prospectar","Lead marcado como NÃO PROSPECTAR")
        else: raise HTTPException(400,"Ação em massa inválida")
    await db.commit(); return {"ok":True,"afetados":len(leads)}


@router.get("/relatorios")
async def relatorios(db:AsyncSession=Depends(get_db),user:Envoxer=Depends(get_current_envoxer)):
    q=select(ComercialLead).where(ComercialLead.deleted_at.is_(None))
    if user.permissao=="envoxer": q=q.where(or_(ComercialLead.responsavel_envoxer_id==user.id,ComercialLead.responsavel_envoxer_id.is_(None)))
    leads=(await db.execute(q)).scalars().all(); ids=[x.id for x in leads]
    msgs=(await db.execute(select(ComercialMessage).where(ComercialMessage.lead_id.in_(ids)).order_by(ComercialMessage.lead_id,ComercialMessage.enviado_em))).scalars().all() if ids else []
    by_channel=defaultdict(lambda:{"enviadas":0,"recebidas":0,"leads_abordados":set(),"leads_responderam":set()})
    for m in msgs:
        a=by_channel[m.canal]; a["enviadas" if m.direcao=="enviada" else "recebidas"]+=1
        a["leads_abordados" if m.direcao=="enviada" else "leads_responderam"].add(m.lead_id)
    channels=[]
    for name,a in by_channel.items():
        channels.append({"canal":name,"enviadas":a["enviadas"],"recebidas":a["recebidas"],"leads_abordados":len(a["leads_abordados"]),"leads_responderam":len(a["leads_responderam"]),"taxa_resposta":round(len(a["leads_responderam"])/len(a["leads_abordados"])*100,1) if a["leads_abordados"] else 0})
    resp_ids={m.lead_id for m in msgs if m.direcao=="recebida"}
    def group(field):
        g=defaultdict(lambda:{"leads":0,"responderam":0,"reunioes":0,"ganhos":0})
        for l in leads:
            key=getattr(l,field) or "Não informado"; a=g[key]; a["leads"]+=1; a["responderam"]+=l.id in resp_ids; a["reunioes"]+=l.status_codigo in STATUS_REUNIAO_OU_DEPOIS; a["ganhos"]+=l.status_codigo=="ganho"
        return [{field:k,**v,"taxa_resposta":round(v["responderam"]/v["leads"]*100,1) if v["leads"] else 0} for k,v in sorted(g.items(),key=lambda kv:kv[1]["leads"],reverse=True)]
    respmap=await _responsaveis(db); gr=defaultdict(lambda:{"leads":0,"responderam":0,"ganhos":0})
    for l in leads:
        key=respmap.get(l.responsavel_envoxer_id).nome if respmap.get(l.responsavel_envoxer_id) else "Sem responsável"; a=gr[key]; a["leads"]+=1; a["responderam"]+=l.id in resp_ids; a["ganhos"]+=l.status_codigo=="ganho"
    golds=[]
    if ids:
        ins=(await db.execute(select(ComercialInsight).where(ComercialInsight.lead_id.in_(ids),ComercialInsight.utilizado.is_(True)))).scalars().all()
        gc=defaultdict(lambda:{"utilizados":0,"leads":set()})
        for x in ins: gc[x.categoria or "outro"]["utilizados"]+=1; gc[x.categoria or "outro"]["leads"].add(x.lead_id)
        golds=[{"categoria":k,"utilizados":v["utilizados"],"leads":len(v["leads"]),"leads_que_responderam":len(v["leads"] & resp_ids),"taxa_resposta":round(len(v["leads"] & resp_ids)/len(v["leads"])*100,1) if v["leads"] else 0} for k,v in gc.items()]
    sent_before_response=[]
    grouped=defaultdict(list)
    for m in msgs: grouped[m.lead_id].append(m)
    for lid,arr in grouped.items():
        count=0
        for m in arr:
            if m.direcao=="recebida": sent_before_response.append(count); break
            if m.direcao=="enviada": count+=1
    metrics={"media_contatos_ate_resposta":round(sum(sent_before_response)/len(sent_before_response),1) if sent_before_response else 0,"taxa_reuniao":round(sum(l.status_codigo in STATUS_REUNIAO_OU_DEPOIS for l in leads)/len(leads)*100,1) if leads else 0,"taxa_proposta":round(sum(l.status_codigo in {"proposta","negociacao","ganho"} for l in leads)/len(leads)*100,1) if leads else 0,"taxa_fechamento":round(sum(l.status_codigo=="ganho" for l in leads)/len(leads)*100,1) if leads else 0}
    return {"metricas":metrics,"por_canal":channels,"por_segmento":group("segmento"),"por_cidade":group("cidade"),"por_responsavel":[{"responsavel":k,**v,"taxa_resposta":round(v["responderam"]/v["leads"]*100,1) if v["leads"] else 0} for k,v in gr.items()],"por_ouro":golds}


@router.get("/notificacoes")
async def notificacoes(db:AsyncSession=Depends(get_db),user:Envoxer=Depends(get_current_envoxer)):
    data=await hoje(db,user); notices=[]
    if data["total"]: notices.append({"tipo":"hoje","titulo":f"Você possui {data['total']} ações comerciais pedindo atenção.","prioridade":"alta"})
    for section in data["sections"]:
        if section["key"]=="respostas" and section["items"]: notices.append({"tipo":"respostas","titulo":f"{len(section['items'])} lead(s) responderam e aguardam ação.","prioridade":"alta"})
        if section["key"]=="analisados" and section["items"]: notices.append({"tipo":"analisados","titulo":f"Existem {len(section['items'])} lead(s) analisados sem abordagem.","prioridade":"media"})
    if data["tarefas_vencidas"]: notices.append({"tipo":"tarefas","titulo":f"{len(data['tarefas_vencidas'])} tarefa(s) comercial(is) estão vencidas.","prioridade":"alta"})
    return notices


@router.get("/config")
async def config_comercial(db:AsyncSession=Depends(get_db),_:Envoxer=Depends(get_current_envoxer)):
    integrations=[d(x) for x in (await db.execute(select(ComercialIntegration).order_by(ComercialIntegration.nome))).scalars().all()]
    for x in integrations:
        if x["provider"]=="openai": x["configurado"]=bool(settings.OPENAI_API_KEY); x["ativo"]=bool(settings.OPENAI_API_KEY); x["config_publica"]={"modelo":settings.OPENAI_MODEL}
    suppressions=[]
    rows=(await db.execute(select(ComercialSuppression).order_by(ComercialSuppression.created_at.desc()).limit(500))).scalars().all()
    if rows:
        lead_ids={x.lead_id for x in rows if x.lead_id}; names={x.id:x.nome_estabelecimento for x in (await db.execute(select(ComercialLead).where(ComercialLead.id.in_(lead_ids)))).scalars().all()} if lead_ids else {}
        suppressions=[{**d(x),"lead_nome":names.get(x.lead_id)} for x in rows]
    tags=[d(x) for x in (await db.execute(select(ComercialTag).order_by(ComercialTag.nome))).scalars().all()]
    return {"integracoes":integrations,"suppression_list":suppressions,"tags":tags,"openai_model":settings.OPENAI_MODEL}


@router.delete("/suppression/{item_id}",status_code=204)
async def remover_suppression(item_id:int,db:AsyncSession=Depends(get_db),_:Envoxer=Depends(get_current_gestor_ou_admin)):
    x=(await db.execute(select(ComercialSuppression).where(ComercialSuppression.id==item_id))).scalar_one_or_none()
    if not x: raise HTTPException(404,"Item não encontrado")
    await db.delete(x); await db.commit()
